"""Format laporan Firebase dengan pola DOCX acuan dan tipografi yang diminta."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import json
import math
import re
import subprocess

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
TMP = ROOT / 'tmp/firebase-report'
REFERENCE = ROOT / 'tmp/doc-format/reference.docx'
REFERENCE_HASH = '8271caeef386dc69e3a84b6c5014f7febc71c87a8e517dc702b2a7e641552208'
assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_HASH
assert (TMP / 'artifact.md').exists()
OUT = ROOT / 'output/docx/Inventora_Tugas_Basis_Data_Firebase.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
DATA_SOURCE = (HERE / 'contoh-data.mjs').read_text(encoding='utf-8')
REPORT = json.loads((HERE / 'hasil-uji.json').read_text(encoding='utf-8'))
RULES = (ROOT / 'firestore.rules').read_text(encoding='utf-8')
assert sha256((HERE / 'contoh-data.mjs').read_bytes()).hexdigest() == REPORT['sha256_data']
assert sha256((ROOT / 'firestore.rules').read_bytes()).hexdigest() == REPORT['sha256_rules']
PAGE_MAP = json.loads((TMP / 'pages.json').read_text(encoding='utf-8-sig')) if (TMP / 'pages.json').exists() else {}
DOC = Document(REFERENCE)
TABLE_PATTERN = deepcopy(DOC.tables[0]._element)
LOGO_P = next(deepcopy(p._p) for p in DOC.paragraphs if p._p.xpath('.//wp:anchor'))
MEMBERS = [p.text for p in DOC.paragraphs if p.style.name == 'List Paragraph'][:5]
BODY = DOC._element.body
SECT = deepcopy(BODY.sectPr)
for child in list(BODY):
    BODY.remove(child)
BODY.append(SECT)
BLACK = RGBColor(0, 0, 0)
WIDTH = DOC.sections[0].page_width - DOC.sections[0].left_margin - DOC.sections[0].right_margin
WIDTH_PT = WIDTH / 12700


def font(run, family='Times New Roman', bold=None):
    run.font.name = family
    run.font.size = Pt(12)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    prop = run._r.get_or_add_rPr()
    rf = prop.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        prop.insert(0, rf)
    for k in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
        rf.set(qn('w:' + k), family)
    for k in ['asciiTheme', 'hAnsiTheme', 'cstheme', 'eastAsiaTheme']:
        rf.attrib.pop(qn('w:' + k), None)
    for k in ['w:szCs']:
        el = prop.find(qn(k))
        if el is None:
            el = OxmlElement(k)
            prop.append(el)
        el.set(qn('w:val'), '24')


for style in DOC.styles:
    if style.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
        continue
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = BLACK
    rp = style.element.get_or_add_rPr()
    szcs = rp.find(qn('w:szCs'))
    if szcs is None:
        szcs = OxmlElement('w:szCs')
        rp.append(szcs)
    szcs.set(qn('w:val'), '24')
    rf = rp.find(qn('w:rFonts'))
    if rf is not None:
        for k in list(rf.attrib):
            if 'Theme' in k or 'theme' in k:
                del rf.attrib[k]
        for k in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            rf.set(qn('w:' + k), 'Times New Roman')
    if style.type == WD_STYLE_TYPE.PARAGRAPH:
        style.paragraph_format.line_spacing = 1.5
DOC.styles['Normal'].paragraph_format.space_after = Pt(6)
DOC.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for border in DOC.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
    border.getparent().remove(border)
for name, before in [('Heading 1', 10), ('Heading 2', 7), ('Heading 3', 6)]:
    s = DOC.styles[name]
    s.font.bold = True
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
if 'Code Block' not in DOC.styles:
    DOC.styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
cs = DOC.styles['Code Block']
cs.base_style = DOC.styles['Normal']
cs.font.name = 'Consolas'
cs.font.size = Pt(12)
cs.font.color.rgb = BLACK
cs.paragraph_format.line_spacing = 1.5
cs.paragraph_format.space_before = Pt(0)
cs.paragraph_format.space_after = Pt(0)
cs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
cs.paragraph_format.keep_together = False
cs.paragraph_format.keep_with_next = False


def p(text='', style=None, bold=False, align=None, before=None, after=None, keep=False):
    para = DOC.add_paragraph(style=style)
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.widow_control = True
    para.paragraph_format.keep_with_next = keep
    if align is not None:
        para.alignment = align
    if before is not None:
        para.paragraph_format.space_before = Pt(before)
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    if text:
        font(para.add_run(text), bold=bold)
    return para


def heading(text, level=1, bookmark=None, newpage=False):
    para = p(text, 'Heading ' + str(level), bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)
    if newpage:
        para.paragraph_format.page_break_before = True
    if bookmark:
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(len(BOOKMARKS) + 100))
        start.set(qn('w:name'), bookmark)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(len(BOOKMARKS) + 100))
        para._p.insert(1, start)
        para._p.append(end)
        BOOKMARKS.append(bookmark)
    return para


def table(headers, rows, widths):
    t = DOC.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    for row_i, values in enumerate([headers] + rows):
        row = t.rows[0] if row_i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if row_i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for i, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'EEEEEE' if row_i == 0 else 'FFFFFF')
            tcpr.insert_element_before(shd, 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign')
            margin = OxmlElement('w:tcMar')
            for side, value in [('top', 75), ('bottom', 75), ('left', 90), ('right', 90)]:
                e = OxmlElement('w:' + side)
                e.set(qn('w:w'), str(value))
                e.set(qn('w:type'), 'dxa')
                margin.append(e)
            tcpr.insert_element_before(margin, 'w:textDirection', 'w:tcFitText', 'w:vAlign')
            cell_borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + side)
                e.set(qn('w:val'), 'single')
                e.set(qn('w:sz'), '8')
                e.set(qn('w:color'), '000000')
                cell_borders.append(e)
            tcpr.insert_element_before(cell_borders, 'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign')
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.keep_with_next = row_i == 0
            font(para.add_run(str(text)), bold=row_i == 0)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '8')
        e.set(qn('w:color'), '000000')
        borders.append(e)
    t._tbl.tblPr.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook')
    p('', after=0)
    return t


def canonical_code(source):
    result, quote, escaped = [], None, False
    i = 0
    while i < len(source):
        ch = source[i]
        if quote:
            result.append(ch)
            if escaped:
                escaped = False
            elif ch == '\\':
                escaped = True
            elif ch == quote:
                quote = None
        elif source[i:i + 2] == '//':
            end = source.find('\n', i)
            i = len(source) if end < 0 else end
            continue
        elif ch in "'\"`":
            quote = ch
            result.append(ch)
        elif not ch.isspace():
            result.append(ch)
        i += 1
    return ''.join(result)


def reflow(source):
    import textwrap
    limit = int(WIDTH_PT // 7.2) - 1
    result = []
    for original in source.splitlines():
        line = original.rstrip()
        if line.lstrip().startswith('//'):
            result.extend(textwrap.wrap(line.lstrip()[2:].strip(),
                          width=limit, initial_indent='// ',
                          subsequent_indent='// '))
            continue
        indent = len(line) - len(line.lstrip())
        while len(line) > limit:
            quote, escaped, cuts = None, False, []
            current_indent = len(line) - len(line.lstrip())
            for i, ch in enumerate(line[:limit]):
                if quote:
                    if escaped:
                        escaped = False
                    elif ch == '\\':
                        escaped = True
                    elif ch == quote:
                        quote = None
                elif ch in "'\"`":
                    quote = ch
                elif i > current_indent and (ch.isspace() or ch in ',()'):
                    cuts.append(i if ch.isspace() else i + 1)
            if not cuts:
                raise ValueError('Baris kode terlalu panjang: ' + line)
            cut = max(cuts)
            result.append(line[:cut].rstrip())
            line = ' ' * (indent + 2) + line[cut:].lstrip()
        result.append(line)
    formatted = '\n'.join(result)
    assert canonical_code(source) == canonical_code(formatted)
    return formatted


CODE_BLOCKS = []
def code(source):
    formatted = reflow(source)
    CODE_BLOCKS.append({'source': source, 'formatted': formatted})
    lines = formatted.splitlines()
    chunks = [lines] if len(lines) <= 13 else [lines[:4], lines[4:-3], lines[-3:]]
    for i, chunk in enumerate(chunks):
        if not chunk:
            continue
        para = p(style='Code Block', align=WD_ALIGN_PARAGRAPH.LEFT,
                 before=0, after=6 if i == len(chunks) - 1 else 0)
        para.paragraph_format.keep_together = len(lines) <= 13 or i in [0, len(chunks) - 1]
        font(para.add_run('\n'.join(chunk)), family='Consolas', bold=False)
    return para


BOOKMARKS = []
CHAPTERS = [
    ('1. Pendahuluan', 'bab1'),
    ('2. Arsitektur Firebase Inventora', 'bab2'),
    ('3. Modul dan penyimpanan aktual', 'bab3'),
    ('4. Struktur koleksi dan dokumen', 'bab4'),
    ('5. Lima dokumen contoh', 'bab5'),
    ('6. Operasi Firebase SDK dan alur PO', 'bab6'),
    ('7. Keamanan dan integritas data', 'bab7'),
    ('8. Hasil pengujian dan kesesuaian proyek', 'bab8'),
    ('Lampiran A. Eksekusi dan sumber proyek', 'lampiranA'),
]

p('Inventora Sistem Enterprise Resource Planning\ndan Pengadaan', 'Title', bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=8)
p('TUGAS KELOMPOK BASIS DATA', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
p('Perancangan dan Implementasi Basis Data\nFirebase Cloud Firestore', bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
DOC._element.body.insert(len(DOC._element.body) - 1, LOGO_P)
from docx.text.paragraph import Paragraph as DocxParagraph
logo_para = DocxParagraph(LOGO_P, DOC._body)
logo_para.paragraph_format.line_spacing = 1.5
logo_para.paragraph_format.space_after = Pt(174)
logo_para.paragraph_format.space_before = Pt(0)
p('Anggota Kelompok', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
for i, member in enumerate(MEMBERS, 1):
    para = p(f'{i}.  {member}', bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    para.paragraph_format.left_indent = Cm(3.2)
p('FAKULTAS INFORMATIKA', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=55, after=6)
p('TELKOM UNIVERSITY PURWOKERTO', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
p('2026', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

heading('Daftar isi', newpage=True)
for title, mark in CHAPTERS:
    para = p(title + '\t' + str(PAGE_MAP.get(mark, '')),
             align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
    para.paragraph_format.tab_stops.add_tab_stop(WIDTH,
        WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

heading(CHAPTERS[0][0], bookmark='bab1', newpage=True)
p('Inventora menggunakan Firebase Cloud Firestore sebagai basis data dokumen dan Firebase Authentication '
  'untuk autentikasi pengguna. Laporan ini menjelaskan implementasi yang terkonfirmasi pada kode proyek: '
  'koleksi purchaseOrders dengan detail barang, penerimaan, dan pembayaran yang tertanam dalam dokumen PO. '
  'Lima dokumen contoh disertakan bersama operasi Firebase SDK dan hasil pengujian pada emulator lokal.')
heading('1.1 Tujuan', 2)
p('Tujuan laporan adalah mendokumentasikan struktur data Inventora secara akurat, memperlihatkan cara '
  'membentuk koleksi melalui penyimpanan dokumen, serta menjelaskan pembacaan dan perubahan data sesuai '
  'alur pengadaan. Contoh dapat dipelajari tanpa mengubah basis data cloud proyek.')
heading('1.2 Ruang lingkup', 2)
p('Cakupan utama adalah pembuatan PO, persetujuan atau penolakan, penerimaan barang, pencocokan penerimaan, '
  'dan pencatatan pembayaran simulasi. Data contoh terdiri atas lima dokumen dalam satu koleksi aktif. '
  'Modul inventori, SDM, payroll, dan keuangan lain tetap dibahas sesuai lokasi penyimpanannya saat ini; '
  'modul tersebut tidak diklaim sudah mempunyai koleksi Firestore.')
heading('1.3 Istilah pada basis data dokumen', 2)
table(['Istilah', 'Pemakaian dalam Inventora'], [
    ['Koleksi', 'Wadah dokumen, yaitu purchaseOrders.'],
    ['Dokumen', 'Satu PO dengan identitas unik pada path purchaseOrders/{poId}.'],
    ['Field', 'Atribut dokumen, misalnya vendor, total, dan status.'],
    ['Map dan array', 'grn dan paymentInfo berupa map; items dan lineItems berupa array map.'],
    ['Security Rules', 'Aturan akses berdasarkan autentikasi, peran, nominal, dan transisi status.'],
], [4.1, 12.6])
p('Firestore menggunakan dokumen dan koleksi, sehingga pembuatan struktur dan pengisian data pada laporan '
  'ini ditunjukkan melalui Firebase SDK, bukan sintaks relasional. Koleksi terbentuk ketika dokumen '
  'pertama berhasil disimpan; tidak diperlukan deklarasi tabel terlebih dahulu.')

heading(CHAPTERS[1][0], bookmark='bab2')
heading('2.1 Komponen utama', 2)
table(['Komponen', 'Peran yang ditemukan pada proyek'], [
    ['Antarmuka React', 'Menyusun PO, menampilkan daftar, menerima perubahan melalui listener, dan mengirim tindakan pengguna.'],
    ['Firebase Authentication', 'Login Google; identitas akun dan custom claim role menentukan hak akses server.'],
    ['Cloud Firestore', 'Penyimpanan dokumen purchaseOrders dan transaksi perubahan status.'],
    ['Firestore Security Rules', 'Memeriksa setiap operasi client; akses koleksi lain ditolak secara bawaan.'],
    ['Penyimpanan browser', 'Data demo, cache PO, inventori, SDM, payroll, dan beberapa modul pendukung.'],
], [4.5, 12.2])
heading('2.2 Alur data', 2)
p('Pada sesi non-demo, PurchaseUnitView membuat dokumen PO menggunakan runTransaction. ProcurementView '
  'membaca koleksi purchaseOrders melalui onSnapshot. Perubahan persetujuan, penerimaan, dan pembayaran '
  'ditulis melalui persistPoTransition. Helper tersebut membaca versi dokumen terkini, memeriksa peran '
  'dan transisi, kemudian mengirim patch dalam transaksi.')
p('Pada mode demo, PO memakai data browser dan tidak ditulis ke Firestore. Cache browser tidak menjadi '
  'bukti bahwa suatu dokumen ada di cloud. Laporan dan contoh pengujian membedakan sesi demo, emulator '
  'lokal, dan implementasi cloud.')
heading('2.3 Bentuk identitas dan relasi', 2)
p('Document ID menjadi identitas utama PO. Field id disimpan pula di dalam dokumen dan harus sama '
  'dengan poId pada path ketika dokumen dibuat. Relasi penerimaan menggunakan grn.poId yang menunjuk '
  'kembali ke PO induknya. items, grn.items, dan paymentInfo disimpan bersama dokumen PO agar data '
  'transaksi dapat dibaca sebagai satu kesatuan.')

heading(CHAPTERS[2][0], bookmark='bab3')
heading('3.1 Peta penyimpanan berdasarkan kode', 2)
table(['Data atau modul', 'Penyimpanan aktual'], [
    ['Purchase order non-demo', 'Koleksi Firestore purchaseOrders; detail item, GRN, dan pembayaran tertanam.'],
    ['Akun dan peran server', 'Firebase Authentication serta custom claim role; bukan koleksi pengguna yang dibuat aplikasi.'],
    ['Inventori', 'localStorage, key inventora_inventory_v1.'],
    ['Pegawai dan payroll', 'localStorage, key inventora_hr_employees_v2 dan inventora_hr_payroll_history_v2.'],
    ['Invoice keuangan', 'localStorage, key inventora_finance_invoices_v1.'],
    ['Langganan dan billing', 'localStorage; antara lain inventora_subscription_v2 dan inventora_subscription_invoices_v1.'],
    ['Anggaran departemen', 'Array DEFAULT_DEPARTMENT_BUDGETS pada kode komponen.'],
    ['Notifikasi dan riwayat transaksi UI', 'Penyimpanan browser; pembayaran utama PO juga tercatat dalam map paymentInfo.'],
], [5.0, 11.7])
heading('3.2 Status blueprint', 2)
p('firebase-blueprint.json mencantumkan purchaseOrders dan activities. Akan tetapi, pemanggilan SDK '
  'yang aktif dan Security Rules proyek hanya mengizinkan purchaseOrders. Tidak ditemukan penulisan '
  'koleksi activities pada implementasi yang diperiksa; aksesnya ditolak oleh aturan bawaan. '
  'Blueprint tersebut diperlakukan sebagai rancangan, bukan bukti koleksi aktif.')
heading('3.3 Kompleksitas yang sudah tersedia', 2)
p('Kompleksitas Inventora terletak pada model dokumen bertingkat, empat peran pengguna, pengendalian '
  'perubahan status, pencocokan penerimaan, transaksi saat menyimpan, serta pencatatan pembayaran '
  'simulasi. Jumlah koleksi tidak mewakili seluruh kompleksitas aplikasi. Penambahan koleksi untuk '
  'modul lokal memerlukan integrasi dan aturan akses tersendiri sebelum dinyatakan terimplementasi.')

heading(CHAPTERS[3][0], bookmark='bab4', newpage=True)
heading('4.1 Hierarki dokumen', 2)
code('''purchaseOrders/{poId}
  id, vendor, desc, date, deliveryDate
  department, subtotal, discountAmount
  taxAmount, withholdingTaxAmount, total
  status, color, threeWayMatchStatus
  items[]
    id, sku, name, qty, unitPrice, subtotal
  lineItems[]
    id, sku, name, qty, unitPrice, totalPrice
  grn (opsional)
    id, grnNumber, poId, receivedDate, receivedBy
    suratJalanNumber, status, inspectorNotes
    items[]
      sku, name, orderedQty, receivedQty
      condition, notes
  paymentInfo (opsional)
    paidAt, method, txId, reference, amount
    accountOrCard, receiptNo, payerName
  lastMutationId (opsional, metadata sinkronisasi)''')
p('grn dan paymentInfo adalah map, bukan subkoleksi. items dan lineItems adalah array map. '
  'lastMutationId merupakan metadata operasional pada helper antrean persetujuan; field tersebut '
  'tidak terdapat pada interface PurchaseOrder utama. Field grn dan paymentInfo tidak disertakan '
  'ketika PO baru dibuat.')
heading('4.2 Field dokumen PO', 2)
table(['Field', 'Tipe', 'Makna atau ketentuan'], [
    ['id', 'string', 'Nomor PO; sama dengan document ID saat create.'],
    ['vendor', 'string', 'Nama pemasok sebagai nilai pada dokumen.'],
    ['desc', 'string', 'Deskripsi atau ringkasan pengadaan.'],
    ['date', 'string', 'Tanggal tampilan yang dibuat aplikasi; bukan Timestamp Firestore.'],
    ['deliveryDate', 'string', 'Opsional pada tipe; formulir mengisi tanggal YYYY-MM-DD.'],
    ['department', 'string', 'Opsional; salah satu dari empat departemen pada tipe proyek.'],
    ['subtotal', 'number', 'Jumlah nilai item; diisi formulir ketika membuat PO.'],
    ['discountAmount', 'number', 'Nominal diskon.'],
    ['taxAmount', 'number', 'Nominal pajak yang disimpan.'],
    ['withholdingTaxAmount', 'number', 'Nominal potongan pajak.'],
    ['total', 'number', 'subtotal - diskon + pajak - potongan.'],
    ['status', 'string', 'Status alur pengadaan sesuai aturan transisi.'],
    ['color', 'string', 'Warna badge antarmuka.'],
    ['threeWayMatchStatus', 'string', 'MATCHED, DISCREPANCY, PENDING_RECEIPT, atau PENDING_INVOICE pada tipe.'],
    ['items', 'array map', 'Detail item utama; pada tipe bersifat opsional.'],
    ['lineItems', 'array map', 'Format kompatibilitas yang juga ditulis formulir; tipe proyek masih any[].'],
], [4.8, 2.6, 9.3])
heading('4.3 Detail item dan penerimaan', 2)
p('Setiap items[] memuat id, sku opsional, name, qty, unitPrice, dan subtotal. '
  'lineItems[] menyalin nilai yang sama tetapi menggunakan totalPrice sebagai nilai baris. '
  'Relasi ke katalog berupa SKU string; proyek belum menyimpan DocumentReference ke koleksi produk.')
table(['Bagian GRN', 'Isi yang dipakai proyek'], [
    ['Identitas', 'id, grnNumber, dan poId berupa string. grn.poId harus sesuai dengan path PO.'],
    ['Penerimaan', 'receivedDate, receivedBy, dan suratJalanNumber berupa string. inspectorNotes opsional.'],
    ['Status', 'MATCHED, DISCREPANCY, atau PENDING pada tipe GoodsReceiptNote.'],
    ['grn.items[]', 'sku opsional, name, orderedQty, receivedQty, condition, dan notes opsional.'],
    ['condition', 'EXCELLENT, GOOD, DAMAGED, atau SHORTAGE.'],
], [4.2, 12.5])
heading('4.4 Informasi pembayaran', 2)
p('paymentInfo menyimpan paidAt, method, txId, reference, accountOrCard, receiptNo, dan payerName '
  'sebagai string serta amount sebagai number. paidAt berisi string waktu ISO pada contoh. amount '
  'harus sama dengan total PO ketika pembayaran dicatat. Pembayaran pada aplikasi ini merupakan '
  'simulasi, bukan konfirmasi transfer dari bank.')
heading('4.5 Konsekuensi model dokumen', 2)
p('Satu PO dapat memuat banyak item, satu map GRN terkini, dan satu map pembayaran. Model saat ini '
  'belum menyediakan koleksi riwayat penerimaan atau pembayaran berulang. Vendor dan departemen '
  'disimpan sebagai string, sehingga tidak ada foreign key atau JOIN yang ditegakkan Firestore. '
  'Konsistensi lintas field bergantung pada transaksi, validasi aplikasi, dan Security Rules.')

heading(CHAPTERS[4][0], bookmark='bab5', newpage=True)
heading('5.1 Ringkasan data akhir', 2)
p('Lima PO berikut bersifat fiktif dan mengikuti nama field aplikasi. Semua PO mula-mula dibuat '
  'berstatus Pending Approval. Status akhirnya dibentuk melalui tindakan yang diizinkan rules, '
  'sehingga contoh penerimaan dan pembayaran tidak melewati tahapan persetujuan.')
final = json.loads((HERE / 'contoh-dokumen.json').read_text(encoding='utf-8'))['documents']
def money(n):
    return 'Rp' + f'{n:,.0f}'.replace(',', '.')
table(['Document ID', 'Status akhir', 'Total'], [
    [po['id'], po['status'], money(po['total'])] for po in final
], [7.0, 5.5, 4.2])
p('Total nilai kelima PO adalah ' + money(REPORT['totalPo']) + '. Angka tersebut mencakup satu PO '
  'yang ditolak dan bukan realisasi belanja. Hanya PO keempat yang mempunyai paymentInfo, '
  'dengan pencatatan pembayaran simulasi sebesar ' + money(REPORT['totalPaid']) + '.')
heading('5.2 Pembentuk lima dokumen', 2)
p('Kode JavaScript berikut berasal dari contoh-data.mjs dan membentuk seluruh field dokumen, '
  'termasuk array detail, map GRN, dan map pembayaran. Variabel pendingDocuments berisi data '
  'yang dapat dibuat menurut rules. finalDocuments adalah hasil akhir yang diharapkan setelah '
  'alur dijalankan. Data akhir lengkap juga tersedia dalam contoh-dokumen.json.')
code(DATA_SOURCE)
heading('5.3 Tanggal dan nominal contoh', 2)
p('Tanggal penerimaan dan pembayaran pada fixture adalah tanggal skenario, bukan waktu transaksi '
  'nyata. Nominal menggunakan bilangan bulat rupiah agar contoh tidak bergantung pada pembulatan '
  'pecahan. Tarif 11 persen hanya mengikuti pengaturan contoh pada kode formulir; laporan ini '
  'tidak menetapkannya sebagai tarif pajak yang berlaku untuk setiap transaksi.')

heading(CHAPTERS[5][0], bookmark='bab6', newpage=True)
p('Potongan berikut menjelaskan inti eksekusi dari uji-firestore.mjs. Skrip lengkap menyiapkan '
  'client emulator dengan peran yang berbeda, menulis data contoh, memeriksa hasil, dan '
  'menutup koneksi. Identitas mock hanya berlaku di emulator; aplikasi cloud memerlukan login '
  'Firebase Authentication serta custom claim yang disetel oleh administrator tepercaya.')
heading('6.1 Koneksi emulator untuk pengujian', 2)
code('''import { initializeApp } from 'firebase/app';
import {
  getFirestore, connectFirestoreEmulator,
  collection, doc, getDocs, query, where,
  runTransaction, updateDoc,
} from 'firebase/firestore';
import {
  pendingDocuments, receiptPatch, paymentPatch,
} from './contoh-data.mjs';

function makeClient(role) {
  const app = initializeApp({
    projectId: 'demo-inventora-tubes',
    apiKey: 'demo-api-key',
  }, role);
  const db = getFirestore(app);
  connectFirestoreEmulator(db, '127.0.0.1', 8085, {
    mockUserToken: { sub: `demo-${role}`, role },
  });
  return db;
}
const warehouse = makeClient('Warehouse Manager');
const finance = makeClient('Finance Manager');''')
heading('6.2 Membentuk koleksi dan mengisi lima dokumen', 2)
code('''for (const po of pendingDocuments) {
  await runTransaction(warehouse, async tx => {
    const ref = doc(warehouse, 'purchaseOrders', po.id);
    const existing = await tx.get(ref);
    if (existing.exists()) {
      throw new Error('Document ID sudah digunakan.');
    }
    tx.set(ref, po);
  });
}''')
p('tx.set menulis satu dokumen PO. Pada dokumen pertama, koleksi purchaseOrders otomatis muncul. '
  'Pengecekan existing mencegah data contoh menimpa PO yang sudah ada. Kelima transaksi dilakukan '
  'terpisah; kegagalan pada dokumen berikutnya tidak membatalkan dokumen sebelumnya. Untuk '
  'mengulang fixture, gunakan instance emulator tugas yang baru.')
heading('6.3 Persetujuan dan penolakan', 2)
code('''for (const i of [1, 2, 3]) {
  const ref = doc(finance, 'purchaseOrders',
    pendingDocuments[i].id);
  await updateDoc(ref, {
    status: 'Disetujui', color: 'green',
  });
}
await updateDoc(doc(finance, 'purchaseOrders',
  pendingDocuments[4].id), {
  status: 'Ditolak', color: 'red',
});''')
heading('6.4 Penerimaan dan pembayaran simulasi', 2)
code('''for (const i of [2, 3]) {
  const po = pendingDocuments[i];
  await updateDoc(doc(warehouse, 'purchaseOrders', po.id),
    receiptPatch(po));
}
const paidPo = pendingDocuments[3];
await updateDoc(doc(finance, 'purchaseOrders', paidPo.id),
  paymentPatch(paidPo));''')
p('PO pertama dibiarkan menunggu persetujuan. PO kedua berhenti pada Disetujui. PO ketiga '
  'berhenti pada Barang Diterima dan memiliki grn. PO keempat mencapai Selesai serta memiliki '
  'grn dan paymentInfo. PO kelima berstatus Ditolak. Pada implementasi antarmuka, perubahan '
  'tersebut memakai persistPoTransition dengan transaksi pembacaan dan pembaruan dokumen terkini.')
heading('6.5 Membaca dan menyaring dokumen', 2)
code('''const snapshot = await getDocs(
  collection(finance, 'purchaseOrders'));
const orders = snapshot.docs.map(item => ({
  ...item.data(), id: item.id,
}));

const completed = query(
  collection(finance, 'purchaseOrders'),
  where('status', '==', 'Selesai'));
const completedSnapshot = await getDocs(completed);
console.log(orders.length, completedSnapshot.size);''')
p('Hasil fixture adalah lima PO dan satu PO berstatus Selesai. ProcurementView menggunakan '
  'onSnapshot untuk pembaruan daftar secara langsung dan menghentikan listener saat komponen '
  'dilepas. Penghapusan PO melalui SDK ditolak oleh rules saat ini.')

heading(CHAPTERS[6][0], bookmark='bab7', newpage=True)
heading('7.1 Peran dan alur yang diizinkan', 2)
table(['Tindakan', 'Peran', 'Prasyarat utama'], [
    ['Baca PO', 'Super Admin, Finance Manager, Warehouse Manager', 'Pengguna terautentikasi dan memiliki claim role yang sesuai.'],
    ['Buat PO', 'Super Admin atau Warehouse Manager', 'Pending Approval; total cocok dengan komponen; tanpa grn dan paymentInfo.'],
    ['Setujui atau tolak', 'Super Admin atau Finance Manager', 'Berawal dari Pending Approval.'],
    ['Catat penerimaan', 'Super Admin atau Warehouse Manager', 'Berawal dari Disetujui atau Selisih Penerimaan; status GRN dan pencocokan selaras.'],
    ['Catat pembayaran', 'Super Admin atau Finance Manager', 'Barang Diterima, MATCHED, belum dibayar, dan nominal sama dengan total PO.'],
    ['Hapus PO', 'Tidak diizinkan', 'allow delete bernilai false.'],
], [3.5, 5.7, 7.5])
p('HR Manager merupakan peran aplikasi untuk modul SDM, tetapi tidak diizinkan membaca PO oleh '
  'rules saat ini. Matriks izin yang disimpan di browser tidak menggantikan custom claim dan '
  'tidak mengubah Security Rules server.')
heading('7.2 Security Rules aktual', 2)
p('Berikut aturan dari firestore.rules proyek. Pembungkusan baris hanya mengubah whitespace agar '
  'kode muat pada halaman. Tidak ada deployment atau perubahan aturan produksi dalam penyusunan laporan.')
code(RULES)
heading('7.3 Batas validasi yang perlu dipahami', 2)
p('Rules memeriksa rumus total ketika create, membatasi field yang boleh berubah, mengendalikan '
  'urutan status, serta menyamakan nominal pembayaran dengan total PO. Rules belum memvalidasi '
  'seluruh struktur items dan grn.items, jumlah subtotal terhadap semua detail, setiap format '
  'tanggal, maupun seluruh properti wajib pada tipe TypeScript. Pemeriksaan jumlah dan kondisi '
  'penerimaan juga dilakukan pada antarmuka; kesesuaian label MATCHED saja bukan pemeriksaan '
  'server yang lengkap atas seluruh item.')
p('Firestore tidak mempunyai foreign key bawaan. grn.poId diperiksa terhadap path PO oleh rules '
  'penerimaan, sedangkan SKU dan nama pemasok masih berupa string. Transaksi mencegah keputusan '
  'berdasarkan versi dokumen yang usang, tetapi tidak menggantikan validasi skema dan aturan bisnis '
  'yang belum ditulis. Hasil uji berikut berlaku pada skenario yang disebutkan.')

heading(CHAPTERS[7][0], bookmark='bab8', newpage=True)
heading('8.1 Lingkungan pengujian', 2)
p('Pengujian dilakukan pada 21 Oktober 2026 melalui emulator Firestore lokal, project ID '
  'demo-inventora-tubes, endpoint 127.0.0.1:8085. Emulator memuat firestore.rules proyek '
  'yang sama dengan laporan. Client SDK menggunakan identitas mock Warehouse Manager, '
  'Finance Manager, HR Manager, dan client tanpa autentikasi. Tidak ada data cloud produksi '
  'yang dibaca atau ditulis oleh skrip pengujian ini.')
heading('8.2 Hasil pemeriksaan', 2)
table(['No', 'Pemeriksaan', 'Hasil'], [
    [str(i), result['name'], 'Lulus'] for i, result in enumerate(REPORT['checks'], 1)
], [1.0, 13.4, 2.3])
p('Seluruh ' + str(REPORT['passed']) + ' pemeriksaan berhasil. Operasi yang sengaja tidak diizinkan '
  'ditolak dengan permission-denied. Pemeriksaan terakhir memastikan penolakan tersebut tidak '
  'mengubah isi lima dokumen. Bukti terperinci tersedia pada hasil-uji.json dan data akhir '
  'pada contoh-dokumen.json.')
heading('8.3 Kesesuaian dengan proyek Inventora', 2)
table(['Aspek', 'Kesesuaian'], [
    ['Platform', 'Cloud Firestore dan Firebase Authentication, sesuai modul firebase proyek.'],
    ['Struktur', 'Satu koleksi aktif purchaseOrders; map dan array sesuai tipe serta formulir.'],
    ['Data contoh', 'Lima dokumen fiktif, dibuat melalui SDK dan alur status yang diizinkan.'],
    ['Keamanan', 'Rules aktual, peran server, nominal pembayaran, dan pembatasan transisi diuji.'],
    ['Modul lokal', 'Inventori, SDM, payroll, dan modul lain dijelaskan menurut penyimpanan aktualnya.'],
], [4.0, 12.7])
heading('8.4 Kesimpulan teknis', 2)
p('Implementasi basis data Inventora yang terkonfirmasi menggunakan model dokumen Firestore '
  'dengan data pengadaan tertanam. Kompleksitas alur pengadaan, peran, dan transaksi tetap ada '
  'walaupun jumlah koleksi aktif hanya satu. Dokumen ini menyajikan implementasi tersebut '
  'beserta lima data contoh yang sudah diuji, tanpa menyatakan modul lokal sudah tersinkronisasi '
  'ke Firebase atau menyatakan pembayaran simulasi sebagai transaksi bank nyata.')

heading(CHAPTERS[8][0], bookmark='lampiranA', newpage=True)
heading('A.1 Berkas paket Firebase', 2)
table(['Berkas', 'Kegunaan'], [
    ['contoh-data.mjs', 'Pembentuk lima dokumen dan patch penerimaan/pembayaran.'],
    ['contoh-dokumen.json', 'Kelima dokumen akhir lengkap, termasuk field bertingkat.'],
    ['uji-firestore.mjs', 'Skrip emulator: membuat data, menjalankan alur, dan memeriksa 18 skenario.'],
    ['hasil-uji.json', 'Bukti pengujian beserta hash sumber data dan rules.'],
    ['README.md', 'Langkah menjalankan contoh secara lokal.'],
], [5.3, 11.4])
p('Seluruh berkas berada dalam folder database/firestore pada proyek Inventora. Skrip uji '
  'mengunci project ID demo dan alamat loopback port 8085. Jangan mengarahkan identitas mock '
  'emulator ke aplikasi cloud.')
heading('A.2 Menjalankan contoh lokal', 2)
p('Gunakan Node.js, dependensi proyek yang sudah terpasang, Java 21, serta JAR emulator '
  'Firestore yang tersedia pada folder .firebase-emulators. Mulai emulator tugas yang baru '
  'dengan firestore.rules proyek, project ID demo-inventora-tubes, dan port 8085. Petunjuk '
  'PowerShell lengkap tersedia di README paket.')
code('node database/firestore/uji-firestore.mjs')
p('Perintah tersebut menjalankan pengisian dan verifikasi, lalu menghasilkan ulang '
  'contoh-dokumen.json serta hasil-uji.json. Proses berhenti bila dokumen contoh sudah ada, '
  'agar tidak menimpa data. Setelah selesai, hentikan hanya instance emulator tugas.')
heading('A.3 Sumber implementasi yang diperiksa', 2)
table(['Sumber proyek', 'Bukti yang digunakan'], [
    ['src/firebase.ts', 'Inisialisasi Cloud Firestore dan Firebase Authentication.'],
    ['src/types.ts', 'PurchaseOrder, POLineItem, GoodsReceiptNote, dan paymentInfo.'],
    ['src/components/PurchaseUnitView.tsx', 'Pembuatan PO dan bentuk payload items serta lineItems.'],
    ['src/components/ProcurementView.tsx', 'Listener purchaseOrders, cache, dan penyimpanan transisi.'],
    ['src/lib/purchaseOrderStore.ts', 'Transaksi perubahan PO dan pemeriksaan custom claim.'],
    ['src/lib/procurement.ts', 'Validasi transisi, nominal, dan normalisasi data pembayaran.'],
    ['firestore.rules', 'Hak akses dan aturan perubahan dokumen.'],
    ['firebase-blueprint.json', 'Rancangan purchaseOrders dan activities yang dibandingkan dengan implementasi.'],
], [7.4, 9.3])
p('Struktur dokumen ditetapkan berdasarkan kode yang dapat diperiksa. Isi Firestore cloud '
  'tidak diekspor; seluruh contoh pada laporan adalah data sintetis. Logo dan identitas '
  'kelompok mengikuti dokumen referensi yang diberikan.')


# Header dan footer mengikuti posisi dan pola acuan.
for section in DOC.sections:
    for para in section.header.paragraphs:
        para.clear()
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.line_spacing = 1.5
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    font(hp.add_run('INVENTORA   /   ENTERPRISE ERP'), bold=True)
    for para in section.footer.paragraphs:
        para.clear()
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.line_spacing = 1.5
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    font(fp.add_run('Firebase Firestore  |  21 Oktober 2026  |  Halaman '), bold=True)
    for tag, val in [('begin', None), ('instruction', 'PAGE'), ('separate', None), ('text', '1'), ('end', None)]:
        run = fp.add_run()
        font(run, bold=True)
        if tag in ['begin', 'separate', 'end']:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tag)
        elif tag == 'instruction':
            e = OxmlElement('w:instrText')
            e.text = val
        else:
            e = OxmlElement('w:t')
            e.text = val
        run._r.append(e)

setting = DOC.settings.element.find(qn('w:updateFields'))
if setting is None:
    setting = OxmlElement('w:updateFields')
    DOC.settings.element.append(setting)
setting.set(qn('w:val'), 'true')
DOC.core_properties.title = 'Inventora Perancangan dan Implementasi Basis Data Firebase Cloud Firestore'
DOC.core_properties.subject = 'Struktur koleksi dan lima dokumen contoh Inventora'

# Default juga dibuat eksplisit untuk teks field atau style turunan.
styles = DOC.styles.element
defaults = styles.find(qn('w:docDefaults'))
if defaults is not None:
    rp = defaults.find('.//' + qn('w:rPr'))
    if rp is not None:
        for child in list(rp):
            if child.tag in [qn('w:rFonts'), qn('w:sz'), qn('w:szCs'), qn('w:color')]:
                rp.remove(child)
        fonts = OxmlElement('w:rFonts')
        for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + k), 'Times New Roman')
        rp.append(fonts)
        for name, value in [('w:sz', '24'), ('w:szCs', '24'), ('w:color', '000000')]:
            e = OxmlElement(name)
            e.set(qn('w:val'), value)
            rp.append(e)

# Normalisasi final semua paragraf dan run aktif, termasuk drawing dan field.
for part in [DOC._element, DOC.sections[0].header._element, DOC.sections[0].footer._element]:
    for pe in part.xpath('.//w:p'):
        prop = pe.get_or_add_pPr()
        spacing = prop.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            prop.append(spacing)
        spacing.set(qn('w:line'), '360')
        spacing.set(qn('w:lineRule'), 'auto')
        code_style = prop.find(qn('w:pStyle'))
        family = 'Consolas' if code_style is not None and code_style.get(qn('w:val')) == cs.style_id else 'Times New Roman'
        from docx.text.run import Run
        for relem in pe.xpath('./w:r'):
            font(Run(relem, None), family=family)

memory = BytesIO()
DOC.save(memory)
memory.seek(0)
editable = {'word/document.xml', 'word/styles.xml', 'word/settings.xml', 'word/header1.xml',
            'word/footer1.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml', 'docProps/core.xml'}
with ZipFile(REFERENCE) as original, ZipFile(memory) as built, ZipFile(OUT, 'w', ZIP_DEFLATED) as final:
    for info in original.infolist():
        payload = built.read(info.filename) if info.filename in editable else original.read(info.filename)
        final.writestr(info, payload)
    for name in set(built.namelist()) - set(original.namelist()):
        final.writestr(name, built.read(name))


with ZipFile(REFERENCE) as original, ZipFile(OUT) as final:
    preserve = set(original.namelist()) - editable
    assert all(original.read(n) == final.read(n) for n in preserve)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for part in ['word/document.xml', 'word/header1.xml', 'word/footer1.xml']:
        x = etree.fromstring(final.read(part))
        for prop in x.xpath('//w:r[w:t]/w:rPr', namespaces=ns):
            assert prop.find(qn('w:sz')).get(qn('w:val')) == '24'
            assert prop.find(qn('w:color')).get(qn('w:val')) == '000000'
        for spacing in x.xpath('//w:p/w:pPr/w:spacing', namespaces=ns):
            assert spacing.get(qn('w:line')) == '360'
            assert spacing.get(qn('w:lineRule')) == 'auto'
    text = ''.join(etree.fromstring(final.read('word/document.xml')).xpath('//w:t/text()', namespaces=ns))
    assert 'CREATE TABLE' not in text and 'INSERT INTO' not in text
    assert all(po['id'] in text for po in json.loads((HERE / 'contoh-dokumen.json').read_text(encoding='utf-8'))['documents'])

(TMP / 'bookmarks.json').write_text(json.dumps(BOOKMARKS, indent=2))
(TMP / 'code-blocks.json').write_text(json.dumps(CODE_BLOCKS, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'docx': str(OUT), 'tables': len(DOC.tables),
    'preserve_parts_unchanged': len(preserve), 'bookmarks': BOOKMARKS}, indent=2))
