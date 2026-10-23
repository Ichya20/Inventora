"""Format laporan MySQL dengan pola DOCX acuan dan tipografi yang diminta."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import json
import math
import re
import subprocess

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
TMP = ROOT / 'tmp/doc-format'
REFERENCE = TMP / 'reference.docx'
REFERENCE_HASH = '8271caeef386dc69e3a84b6c5014f7febc71c87a8e517dc702b2a7e641552208'
assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_HASH
assert (TMP / 'artifact.md').exists()
OUT = ROOT / 'output/docx/Inventora_Tugas_Basis_Data_MySQL.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
SQL = (HERE / 'inventora_tubes.sql').read_text(encoding='utf-8')
REPORT = json.loads((HERE / 'hasil-uji.json').read_text(encoding='utf-8'))
assert sha256(SQL.encode()).hexdigest() == REPORT['sha256_sql']
CREATES = {m[1]: m[0] for m in re.finditer(r'CREATE TABLE (\w+) \([\s\S]+?\) ENGINE=InnoDB;', SQL)}
INSERTS = {m[1]: m[0] for m in re.finditer(r'INSERT INTO (\w+)\b[\s\S]+?;', SQL)}
PAGE_MAP = json.loads((TMP / 'pages.json').read_text()) if (TMP / 'pages.json').exists() else {}
DOC = Document(REFERENCE)
TABLE_PATTERN = deepcopy(DOC.tables[0]._element)
LOGO_P = next(deepcopy(p._p) for p in DOC.paragraphs if p._p.xpath('.//wp:anchor'))
MEMBERS = [p.text for p in DOC.paragraphs if p.style.name == 'List Paragraph'][:5]
BODY = DOC._element.body
SECT = deepcopy(BODY.sectPr)
for child in list(BODY):
    BODY.remove(child)
BODY.append(SECT)
BLACK = RGBColor(0, 0, 0)
WIDTH = DOC.sections[0].page_width - DOC.sections[0].left_margin - DOC.sections[0].right_margin
WIDTH_PT = WIDTH / 12700


def font(run, family='Times New Roman', bold=None):
    run.font.name = family
    run.font.size = Pt(12)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    prop = run._r.get_or_add_rPr()
    rf = prop.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        prop.insert(0, rf)
    for k in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
        rf.set(qn('w:' + k), family)
    for k in ['asciiTheme', 'hAnsiTheme', 'cstheme', 'eastAsiaTheme']:
        rf.attrib.pop(qn('w:' + k), None)
    for k in ['w:szCs']:
        el = prop.find(qn(k))
        if el is None:
            el = OxmlElement(k)
            prop.append(el)
        el.set(qn('w:val'), '24')


for style in DOC.styles:
    if style.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
        continue
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = BLACK
    rp = style.element.get_or_add_rPr()
    szcs = rp.find(qn('w:szCs'))
    if szcs is None:
        szcs = OxmlElement('w:szCs')
        rp.append(szcs)
    szcs.set(qn('w:val'), '24')
    rf = rp.find(qn('w:rFonts'))
    if rf is not None:
        for k in list(rf.attrib):
            if 'Theme' in k or 'theme' in k:
                del rf.attrib[k]
        for k in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            rf.set(qn('w:' + k), 'Times New Roman')
    if style.type == WD_STYLE_TYPE.PARAGRAPH:
        style.paragraph_format.line_spacing = 1.5
DOC.styles['Normal'].paragraph_format.space_after = Pt(6)
DOC.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for border in DOC.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
    border.getparent().remove(border)
for name, before in [('Heading 1', 10), ('Heading 2', 7), ('Heading 3', 6)]:
    s = DOC.styles[name]
    s.font.bold = True
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
if 'SQL Code' not in DOC.styles:
    DOC.styles.add_style('SQL Code', WD_STYLE_TYPE.PARAGRAPH)
cs = DOC.styles['SQL Code']
cs.base_style = DOC.styles['Normal']
cs.font.name = 'Consolas'
cs.font.size = Pt(12)
cs.font.color.rgb = BLACK
cs.paragraph_format.line_spacing = 1.5
cs.paragraph_format.space_before = Pt(0)
cs.paragraph_format.space_after = Pt(0)
cs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
cs.paragraph_format.keep_together = False
cs.paragraph_format.keep_with_next = False


def p(text='', style=None, bold=False, align=None, before=None, after=None, keep=False):
    para = DOC.add_paragraph(style=style)
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.widow_control = True
    para.paragraph_format.keep_with_next = keep
    if align is not None:
        para.alignment = align
    if before is not None:
        para.paragraph_format.space_before = Pt(before)
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    if text:
        font(para.add_run(text), bold=bold)
    return para


def heading(text, level=1, bookmark=None, newpage=False):
    para = p(text, 'Heading ' + str(level), bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)
    if newpage:
        para.paragraph_format.page_break_before = True
    if bookmark:
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(len(BOOKMARKS) + 100))
        start.set(qn('w:name'), bookmark)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(len(BOOKMARKS) + 100))
        para._p.insert(1, start)
        para._p.append(end)
        BOOKMARKS.append(bookmark)
    return para


def table(headers, rows, widths):
    t = DOC.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    for row_i, values in enumerate([headers] + rows):
        row = t.rows[0] if row_i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if row_i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for i, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'EEEEEE' if row_i == 0 else 'FFFFFF')
            tcpr.insert_element_before(shd, 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign')
            margin = OxmlElement('w:tcMar')
            for side, value in [('top', 75), ('bottom', 75), ('left', 90), ('right', 90)]:
                e = OxmlElement('w:' + side)
                e.set(qn('w:w'), str(value))
                e.set(qn('w:type'), 'dxa')
                margin.append(e)
            tcpr.insert_element_before(margin, 'w:textDirection', 'w:tcFitText', 'w:vAlign')
            cell_borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + side)
                e.set(qn('w:val'), 'single')
                e.set(qn('w:sz'), '8')
                e.set(qn('w:color'), '000000')
                cell_borders.append(e)
            tcpr.insert_element_before(cell_borders, 'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign')
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.keep_with_next = row_i == 0
            font(para.add_run(str(text)), bold=row_i == 0)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '8')
        e.set(qn('w:color'), '000000')
        borders.append(e)
    t._tbl.tblPr.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook')
    p('', after=0)
    return t


def normalize_sql(text):
    out = []
    quoted = False
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == "'":
            quoted = not quoted
        if quoted or not ch.isspace():
            out.append(ch)
        i += 1
    return ''.join(out)


def reflow(text):
    limit = int(WIDTH_PT // 7.2) - 1
    output = []
    for original in text.splitlines():
        rest = original.rstrip()
        indent = len(rest) - len(rest.lstrip())
        continuation = ' ' * (indent + 2)
        while len(rest) > limit:
            quoted = False
            breaks = []
            for i, ch in enumerate(rest):
                if ch == "'":
                    quoted = not quoted
                if not quoted and (ch.isspace() or ch == ',') and indent < i < limit:
                    breaks.append(i + 1 if ch == ',' else i)
            if not breaks:
                raise ValueError('Literal atau token terlalu panjang: ' + rest)
            cut = max(breaks)
            output.append(rest[:cut].rstrip())
            rest = continuation + rest[cut:].lstrip()
        output.append(rest)
    result = '\n'.join(output)
    assert normalize_sql(text) == normalize_sql(result)
    return result


def code(text):
    lines = reflow(text).splitlines()
    # Blok singkat tetap utuh; blok panjang dimulai dengan minimal empat baris
    # dan ditutup dengan tiga baris agar tidak meninggalkan judul atau ekor kode.
    chunks = [lines] if len(lines) <= 14 else [lines[:4], lines[4:-3], lines[-3:]]
    for i, chunk in enumerate(chunks):
        if not chunk:
            continue
        para = p(style='SQL Code', align=WD_ALIGN_PARAGRAPH.LEFT, before=0,
                 after=6 if i == len(chunks) - 1 else 0)
        para.paragraph_format.keep_together = len(lines) <= 14 or i in [0, len(chunks) - 1]
        font(para.add_run('\n'.join(chunk)), family='Consolas', bold=False)
    return para


def figure(path, caption):
    para = p(align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
    para.add_run().add_picture(str(path), width=WIDTH)
    p(caption, style='Caption', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=8)


def make_diagrams():
    pdfmetrics.registerFont(TTFont('TNR', 'C:/Windows/Fonts/times.ttf'))
    pdfmetrics.registerFont(TTFont('TNRB', 'C:/Windows/Fonts/timesbd.ttf'))
    poppler = 'C:/Users/ICHYA ULUMIDDIN/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe'
    width = WIDTH_PT

    def draw(kind, height, nodes, edges):
        pdf = TMP / (kind + '.pdf')
        c = canvas.Canvas(str(pdf), pagesize=(width, height))
        c.setFillColorRGB(0, 0, 0)
        c.setStrokeColorRGB(0, 0, 0)
        # Gaps diperlebar agar label 1:0..1 tidak menyentuh batas kotak.
        xmap = {0: 0, 70: 65, 89: 85, 96: 92, 140: 130, 153: 150.5,
                166: 171, 236: 236, 254: 254, 260: 260, 306: 301,
                318: 321.5, 319: 321.5, 332: 342, 402: 407, 420: 425}
        for points, label, lx, ly in edges:
            points = [(xmap.get(x, x), y) for x, y in points]
            lx = xmap.get(lx, lx)
            path = c.beginPath()
            path.moveTo(*points[0])
            for pt in points[1:]:
                path.lineTo(*pt)
            c.drawPath(path)
            x0, y0 = points[-2]
            x1, y1 = points[-1]
            a = math.atan2(y1 - y0, x1 - x0)
            for d in [-0.4, 0.4]:
                c.line(x1, y1, x1 - 6 * math.cos(a + d), y1 - 6 * math.sin(a + d))
            c.setFont('TNR', 12)
            c.drawCentredString(lx, ly, label)
        for x, y, name, key in nodes:
            x = xmap.get(x, x)
            c.rect(x, y, 130, 70)
            c.setFont('TNRB', 12)
            c.drawCentredString(x + 65, y + 48, name)
            c.setFont('TNR', 12)
            for j, line in enumerate(key):
                c.drawCentredString(x + 65, y + 29 - 18 * j, line)
        c.save()
        prefix = TMP / kind
        subprocess.run([poppler, '-r', '180', '-singlefile', '-png', str(pdf), str(prefix)], check=True, capture_output=True)
        return prefix.with_suffix('.png')

    hr = draw('relasi-organisasi', 330, [
        (0, 240, 'departemen', ['PK KodeDepartemen']),
        (166, 240, 'pegawai', ['PK KodePegawai']),
        (332, 240, 'pengguna', ['PK KodePengguna']),
        (0, 50, 'anggaran_departemen', ['PK KodeAnggaran']),
        (166, 50, 'payroll', ['PK KodePayroll']),
    ], [
        ([(140, 275), (166, 275)], '1:N', 153, 290),
        ([(306, 275), (332, 275)], '1:0..1', 319, 290),
        ([(70, 240), (70, 120)], '1:N', 89, 182),
        ([(236, 240), (236, 120)], '1:N', 254, 182),
        ([(402, 240), (402, 85), (306, 85)], '1:N', 420, 182),
    ])
    procurement = draw('relasi-pengadaan', 470, [
        (0, 390, 'kategori_produk', ['PK KodeKategori']),
        (166, 390, 'produk', ['PK KodeProduk']),
        (332, 390, 'supplier', ['PK KodeSupplier']),
        (0, 220, 'detail_purchase_order', ['PK KodePO dan', 'KodeProduk']),
        (166, 220, 'purchase_order', ['PK KodePO']),
        (332, 220, 'persetujuan_po', ['PK KodePersetujuan']),
        (0, 50, 'detail_penerimaan', ['PK KodePO dan', 'KodeProduk']),
        (166, 50, 'penerimaan_barang', ['PK KodePO']),
        (332, 50, 'pembayaran_po', ['PK KodeBayar']),
    ], [
        ([(140, 425), (166, 425)], '1:N', 153, 439),
        ([(236, 390), (236, 360), (70, 360), (70, 290)], '1:N', 89, 318),
        ([(402, 390), (402, 335), (236, 335), (236, 290)], '1:N', 318, 341),
        ([(166, 255), (140, 255)], '1:N', 153, 270),
        ([(306, 255), (332, 255)], '1:0..1', 319, 270),
        ([(236, 220), (236, 120)], '1:0..1', 260, 173),
        ([(70, 220), (70, 120)], '1:0..1', 96, 173),
        ([(166, 85), (140, 85)], '1:N', 153, 100),
        ([(306, 85), (332, 85)], '1:0..1', 319, 100),
    ])
    return hr, procurement


BOOKMARKS = []
CHAPTERS = [
    ('1. Pendahuluan dan dasar perancangan', 'bab1'),
    ('2. Arsitektur basis data', 'bab2'),
    ('3. Rancangan modul', 'bab3'),
    ('4. Rancangan data dan relasi', 'bab4'),
    ('5. Alur transaksi dan eksekusi', 'bab5'),
    ('6. Sintaks SQL dan data contoh', 'bab6'),
    ('7. Integritas data dan pengujian', 'bab7'),
    ('8. Matriks pemenuhan tugas dan keputusan teknis', 'bab8'),
    ('Lampiran A. Query verifikasi', 'lampiranA'),
    ('Lampiran B. Ringkasan data dan berkas', 'lampiranB'),
]

# Sampul memakai komponen dan identitas dari dokumen acuan.
p('Inventora Sistem Enterprise Resource Planning\ndan Pengadaan', 'Title', bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=8)
p('TUGAS KELOMPOK BASIS DATA', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
p('Perancangan dan Implementasi Basis Data MySQL', bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
DOC._element.body.insert(len(DOC._element.body) - 1, LOGO_P)
from docx.text.paragraph import Paragraph as DocxParagraph
logo_para = DocxParagraph(LOGO_P, DOC._body)
logo_para.paragraph_format.line_spacing = 1.5
logo_para.paragraph_format.space_after = Pt(192)
logo_para.paragraph_format.space_before = Pt(0)
p('Anggota Kelompok', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
for i, member in enumerate(MEMBERS, 1):
    para = p(f'{i}.  {member}', bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    para.paragraph_format.left_indent = Cm(3.2)
p('FAKULTAS INFORMATIKA', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=55, after=6)
p('TELKOM UNIVERSITY PURWOKERTO', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
p('2026', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

heading('Daftar isi', newpage=True)
for name, mark in CHAPTERS:
    para = p(name + '\t' + str(PAGE_MAP.get(mark, '')), align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
    para.paragraph_format.tab_stops.add_tab_stop(WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

heading(*CHAPTERS[0][:1], bookmark='bab1', newpage=True)
p('Dokumen ini menyajikan rancangan basis data relasional Inventora menggunakan MySQL. '
  'Cakupan tugas meliputi pembuatan seluruh tabel dan pengisian lima baris data contoh pada setiap tabel. '
  'Rancangan terdiri atas 14 tabel dengan 70 baris data dan 19 constraint foreign key. '
  'Skrip telah dijalankan pada MySQL Community Server 8.0.41 dan seluruh 20 pemeriksaan berhasil.')
heading('1.1 Tujuan dokumen', 2)
p('Tujuan dokumen adalah menjelaskan struktur tabel, hubungan data, sintaks CREATE dan INSERT, '
  'serta hasil verifikasi agar rancangan dapat dipelajari dan dijalankan kembali untuk tugas kelompok.')
heading('1.2 Ruang lingkup', 2)
p('Lingkup sistem mencakup organisasi, pegawai dan akun pengguna, master produk dan supplier, '
  'purchase order, persetujuan, penerimaan barang, pembayaran simulasi, payroll, serta anggaran departemen. '
  'Sistem yang dimodelkan merupakan inventori dan pengadaan internal. Modul penjualan POS tidak dimasukkan '
  'ke dalam paket SQL ini.')
heading('1.3 Konvensi perancangan', 2)
table(['Istilah', 'Pengertian'], [
    ['PK', 'Primary key atau identitas unik suatu baris.'],
    ['FK', 'Foreign key atau referensi ke key pada tabel induk.'],
    ['UNIQUE', 'Pembatas nilai atau gabungan nilai agar tidak berulang.'],
    ['CHECK', 'Pembatas nilai dalam satu baris, misalnya quantity harus positif.'],
    ['Generated column', 'Kolom yang dihitung otomatis dari kolom lain dalam baris yang sama.'],
], [4.1, 12.6])

heading(*CHAPTERS[1][:1], bookmark='bab2')
heading('2.1 Platform dan penyimpanan', 2)
p('Target eksekusi adalah MySQL 8.0.16 atau versi yang lebih baru karena CHECK mulai ditegakkan '
  'pada versi tersebut. Seluruh tabel menggunakan InnoDB dengan karakter utf8mb4 dan collation '
  'utf8mb4_0900_ai_ci. Kolom uang menggunakan DECIMAL(15,2), sedangkan tanggal disimpan sebagai DATE '
  'atau DATETIME sesuai kebutuhan.')
heading('2.2 Hubungan dengan aplikasi Inventora', 2)
p('Paket ini merupakan rancangan relasional untuk tugas basis data. Aplikasi Inventora yang ada '
  'masih menggunakan Firebase. Eksekusi skrip SQL tidak memigrasikan data atau mengubah koneksi '
  'aplikasi secara otomatis. Migrasi aplikasi memerlukan backend, pemetaan data, dan penyesuaian '
  'autentikasi serta otorisasi tersendiri.')

heading(*CHAPTERS[2][:1], bookmark='bab3')
table(['Modul', 'Tabel'], [
    ['Organisasi dan akses', 'departemen, pegawai, pengguna'],
    ['Master inventori', 'supplier, kategori_produk, produk'],
    ['Pengadaan', 'purchase_order, detail_purchase_order, persetujuan_po'],
    ['Penerimaan barang', 'penerimaan_barang, detail_penerimaan'],
    ['Keuangan dan SDM', 'pembayaran_po, payroll, anggaran_departemen'],
], [5.0, 11.7])
heading('3.1 Organisasi dan akun pengguna', 2)
p('Departemen menjadi induk pegawai dan anggaran. Setiap pegawai memiliki paling banyak satu '
  'akun pengguna. Lima akun contoh menggunakan empat peran, yaitu SUPER_ADMIN, WAREHOUSE_MANAGER, '
  'FINANCE_MANAGER, dan HR_MANAGER. AuthUID dipakai sebagai pemetaan identitas eksternal, bukan '
  'sebagai penyimpanan password.')
heading('3.2 Pengadaan dan penerimaan', 2)
p('Satu PO dapat memiliki banyak item produk. Dalam lingkup tugas, setiap PO memiliki paling '
  'banyak satu keputusan persetujuan, satu dokumen penerimaan, dan satu pembayaran penuh. '
  'Penerimaan parsial berulang, cicilan, dan riwayat percobaan pembayaran belum dimodelkan. '
  'FK gabungan pada detail_penerimaan memastikan produk yang diterima terdapat pada PO yang sama.')
heading('3.3 Inventori dan data keuangan', 2)
p('Stok produk merepresentasikan satu lokasi dan merupakan snapshot akhir data contoh. '
  'INSERT penerimaan tidak otomatis menaikkan stok. Harga pada detail PO dan komponen payroll '
  'merupakan snapshot historis. Perubahan master tidak mengubah transaksi yang sudah dicatat. '
  'Sisa anggaran diperoleh melalui query, bukan disimpan ulang sebagai nilai turunan.')

hr, procurement = make_diagrams()
heading(*CHAPTERS[3][:1], bookmark='bab4', newpage=True)
heading('4.1 Relasi organisasi dan payroll', 2)
p('Panah menunjukkan arah dari tabel induk ke tabel anak. Label 1:N berarti satu induk dapat '
  'memiliki banyak anak, sedangkan 1:0..1 berarti paling banyak satu anak. PK menunjukkan primary key.')
figure(hr, 'Gambar 1. Relasi utama organisasi, pengguna, payroll, dan anggaran.')
heading('4.2 Relasi pengadaan dan inventori', 2, newpage=True)
p('Diagram berikut menampilkan relasi utama pengadaan. FK pelaku transaksi dan departemen PO '
  'dirinci bersama seluruh FK pada bagian 4.3.')
figure(procurement, 'Gambar 2. Relasi utama pengadaan, produk, penerimaan, dan pembayaran.')
heading('4.3 Daftar seluruh foreign key', 2)
p('Seluruh FK menggunakan ON UPDATE RESTRICT dan ON DELETE RESTRICT. Perubahan atau '
  'penghapusan key induk ditolak selama masih dirujuk oleh tabel anak.')
relations = []
for name, ddl in CREATES.items():
    for m in re.finditer(r'FOREIGN KEY \(([^)]+)\)\s+REFERENCES (\w+)\(([^)]+)\)', ddl):
        relations.append([name + '\n' + m[1], m[2] + '\n' + m[3]])
assert len(relations) == 19
table(['Tabel anak dan kolom FK', 'Tabel induk dan key referensi'], relations, [8.35, 8.35])
p('Satu FK pada detail_penerimaan merupakan FK gabungan dua kolom, yaitu KodePO dan '
  'KodeProduk. Dengan demikian terdapat 19 constraint FK yang mencakup 20 pasangan kolom referensi.')

heading(*CHAPTERS[4][:1], bookmark='bab5')
heading('5.1 Alur transaksi contoh', 2)
p('Urutan transaksi adalah pembuatan PO dan detail, pencatatan keputusan persetujuan, '
  'penerimaan barang beserta detail, lalu pembayaran simulasi. Kelima PO contoh berstatus SELESAI '
  'dan memiliki persetujuan, penerimaan yang sesuai, serta pembayaran sebesar total PO. '
  'Hak akses dan urutan proses antartabel perlu ditegakkan oleh aplikasi atau prosedur tambahan.')
heading('5.2 Urutan menjalankan skrip', 2)
for n, text in enumerate([
    'Siapkan MySQL 8.0.16+ dan koneksi server. Gunakan schema inventora_tubes yang belum ada atau masih kosong.',
    'Buka inventora_tubes.sql melalui File > Open SQL Script di MySQL Workbench, lalu jalankan seluruh pernyataan.',
    'Jika menyalin dari dokumen, jalankan persiapan schema dan seluruh CREATE TABLE lebih dahulu. Setelah itu jalankan START TRANSACTION, seluruh INSERT sesuai urutan, dan COMMIT.',
    'Buka verifikasi.sql dan jalankan seluruh query. Setiap tabel harus berisi lima baris; query ketidaksesuaian subtotal, pembayaran, serta penerimaan harus kosong.',
], 1):
    p(f'{n}. {text}')
p('Skrip tidak berisi DROP dan tidak dirancang untuk dijalankan berulang pada schema yang '
  'telah terisi. Jangan gunakan --force. Bila terjadi error, hentikan dan periksa error pertama. '
  'DDL MySQL melakukan commit implisit; transaksi INSERT tidak membatalkan CREATE TABLE.')

heading(*CHAPTERS[5][:1], bookmark='bab6', newpage=True)
heading('6.1 Persiapan schema', 2)
prep = SQL[SQL.index('CREATE DATABASE'):SQL.index('-- TABEL 01')].strip()
mode = re.search(r"SET SESSION sql_mode =\s*'([^']+)';", prep)[1]
mode_parts = ['STRICT_TRANS_TABLES,NO_ZERO_IN_DATE,NO_ZERO_DATE,',
              'ERROR_FOR_DIVISION_BY_ZERO,', 'NO_ENGINE_SUBSTITUTION']
assert ''.join(mode_parts) == mode
prep = re.sub(r"SET SESSION sql_mode =\s*'[^']+';",
              'SET SESSION sql_mode = CONCAT(\n' +
              ',\n'.join("  '" + x + "'" for x in mode_parts) + '\n);', prep)
code(prep)
p('CONCAT pada pengaturan sql_mode menggabungkan bagian teks menjadi nilai yang sama '
  'dengan skrip sumber. Penyajian ini menjaga panjang baris agar tetap berada di dalam margin.')
heading('6.2 CREATE seluruh tabel', 2)
for i, (name, ddl) in enumerate(CREATES.items(), 1):
    heading(f'6.2.{i} Tabel {name}', 3)
    code(ddl)
heading('6.3 INSERT lima data pada setiap tabel', 2)
p('Seluruh data berikut bersifat fiktif. INSERT mengikuti urutan dependensi FK. '
  'Kolom generated tidak disertakan pada daftar nilai karena dihitung oleh MySQL.')
code('START TRANSACTION;')
for i, (name, statement) in enumerate(INSERTS.items(), 1):
    heading(f'6.3.{i} Data {name}', 3)
    code(statement)
code('COMMIT;')

heading(*CHAPTERS[6][:1], bookmark='bab7')
heading('7.1 Lingkungan dan metode', 2)
p('Pengujian dilaksanakan pada 21 Oktober 2026 menggunakan MySQL Community Server 8.0.41 '
  'pada instance sementara di 127.0.0.1:33308, terpisah dari data aplikasi. Pengujian mencakup '
  'eksekusi skrip, pemeriksaan data dan relasi, perhitungan nilai, serta delapan percobaan data '
  'tidak valid. Instance sementara telah dihentikan setelah pengujian.')
heading('7.2 Hasil pemeriksaan', 2)
table(['No', 'Pemeriksaan', 'Hasil'],
      [[i, t['uji'], 'Lulus'] for i, t in enumerate(REPORT['uji'], 1)], [1.0, 13.4, 2.3])
p('Pelanggaran FK ditolak dengan ERROR 1452, pelanggaran CHECK dengan ERROR 3819, dan '
  'duplikasi key dengan ERROR 1062. Percobaan dilakukan dalam transaksi yang dibatalkan. '
  'Setelah uji negatif, 14 tabel tetap berisi masing-masing lima baris.')
heading('7.3 Hasil agregat dan batas pembuktian', 2)
table(['Parameter', 'Hasil'], [
    ['Jumlah tabel dan data', '14 tabel, 70 baris; masing-masing 5 baris'],
    ['Total PO dan pembayaran', 'Rp860.970.000,00'],
    ['Total payroll Oktober 2026', 'Rp125.690.000,00'],
    ['Jumlah stok produk contoh', '339 unit'],
    ['Produk di bawah minimum', 'PRD003 - NVMe 8TB; stok 8, minimum 15'],
], [8.2, 8.5])
p('Query tidak menemukan ketidaksesuaian subtotal header-detail, pembayaran, atau isi '
  'penerimaan pada data contoh. Pemeriksaan peran dan urutan tanggal membuktikan konsistensi '
  'data contoh, bukan penegakan otomatis seluruh aturan proses bisnis oleh FK. Nominal pajak, '
  'BPJS, dan PPh21 adalah ilustrasi dan bukan mesin perhitungan perpajakan produksi.')

heading(*CHAPTERS[7][:1], bookmark='bab8')
heading('8.1 Matriks pemenuhan tugas', 2)
table(['Ketentuan', 'Pemenuhan'], [
    ['CREATE semua tabel', '14 pernyataan CREATE TABLE lengkap pada bagian 6.2.'],
    ['INSERT lima data setiap tabel', '14 pernyataan INSERT, total 70 baris pada bagian 6.3.'],
    ['Relasi antartabel', '19 FK termasuk satu FK gabungan, dijelaskan pada bagian 4.'],
    ['Dokumen PDF', 'Laporan berisi rancangan, sintaks, data contoh, dan bukti pengujian.'],
], [6.0, 10.7])
heading('8.2 Keputusan teknis dan batas rancangan', 2)
p('Penggunaan key gabungan pada detail transaksi mencegah produk yang sama berulang '
  'dalam satu dokumen. UNIQUE pada pegawai-periode mencegah payroll bulanan ganda. Generated '
  'columns menghitung nilai turunan dalam satu baris. FK pembayaran mengharuskan keberadaan '
  'dokumen penerimaan sebelum pembayaran dicatat.')
p('Aturan jumlah penerimaan, total pembayaran, kesamaan subtotal terhadap detail, otorisasi '
  'peran, transisi status, dan pembaruan stok memerlukan validasi aplikasi dalam transaksi atau '
  'prosedur/trigger tambahan. Tema inventori dan pengadaan perlu sesuai dengan proposal TUBES '
  'yang disetujui dosen. Kompleksitas rancangan tidak menjamin nilai tertentu.')

heading(*CHAPTERS[8][:1], bookmark='lampiranA', newpage=True)
p('Jalankan query berikut setelah seluruh INSERT selesai. Semua query berasal dari '
  'verifikasi.sql dan telah berhasil dijalankan pada lingkungan pengujian.')
code('USE inventora_tubes;')
verification = (HERE / 'verifikasi.sql').read_text(encoding='utf-8')
for part in re.split(r'(?m)^-- (?=\d+\.)', verification)[1:]:
    comment, body = part.split('\n', 1)
    number, title = comment.split('. ', 1)
    heading(f'A.{number} {title.rstrip(".")}', 2)
    code(body.strip())

heading(*CHAPTERS[9][:1], bookmark='lampiranB', newpage=True)
heading('B.1 Ringkasan anggaran contoh', 2)
p('Nilai berikut dalam rupiah untuk kuartal 4 tahun 2026. Nilai PO mencakup komitmen PO '
  'selain yang ditolak pada periode yang sama, bukan hanya arus kas yang telah dibayar.')
table(['Departemen', 'Anggaran', 'Nilai PO', 'Sisa anggaran'], [
    ['Engineering dan IT', '1.000.000.000', '555.000.000', '445.000.000'],
    ['Operasional', '300.000.000', '159.450.000', '140.550.000'],
    ['Keuangan', '150.000.000', '66.600.000', '83.400.000'],
    ['Sumber Daya Manusia', '100.000.000', '36.630.000', '63.370.000'],
    ['Pergudangan', '120.000.000', '43.290.000', '76.710.000'],
], [5.1, 4.0, 3.8, 3.8])
heading('B.2 Berkas pendukung', 2)
table(['Berkas', 'Kegunaan'], [
    ['inventora_tubes.sql', 'Skrip CREATE dan INSERT siap dieksekusi.'],
    ['verifikasi.sql', 'Delapan bagian query verifikasi.'],
    ['hasil-uji.json', 'Rincian 20 pemeriksaan MySQL.'],
    ['hasil-query.txt', 'Keluaran query dari data contoh.'],
    ['README.md', 'Panduan eksekusi dan batas rancangan.'],
], [6.2, 10.5])
p('Berkas pendukung berada pada folder database/tubes di dalam proyek Inventora.')

# Header dan footer mengikuti posisi dan pola acuan.
for section in DOC.sections:
    for para in section.header.paragraphs:
        para.clear()
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.line_spacing = 1.5
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    font(hp.add_run('INVENTORA   /   ENTERPRISE ERP'), bold=True)
    for para in section.footer.paragraphs:
        para.clear()
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.line_spacing = 1.5
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    font(fp.add_run('SQL MySQL  |  21 Oktober 2026  |  Halaman '), bold=True)
    for tag, val in [('begin', None), ('instruction', 'PAGE'), ('separate', None), ('text', '1'), ('end', None)]:
        run = fp.add_run()
        font(run, bold=True)
        if tag in ['begin', 'separate', 'end']:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tag)
        elif tag == 'instruction':
            e = OxmlElement('w:instrText')
            e.text = val
        else:
            e = OxmlElement('w:t')
            e.text = val
        run._r.append(e)

setting = DOC.settings.element.find(qn('w:updateFields'))
if setting is None:
    setting = OxmlElement('w:updateFields')
    DOC.settings.element.append(setting)
setting.set(qn('w:val'), 'true')
DOC.core_properties.title = 'Inventora Perancangan dan Implementasi Basis Data MySQL'
DOC.core_properties.subject = 'Tugas kelompok CREATE TABLE dan INSERT data'

# Default juga dibuat eksplisit untuk teks field atau style turunan.
styles = DOC.styles.element
defaults = styles.find(qn('w:docDefaults'))
if defaults is not None:
    rp = defaults.find('.//' + qn('w:rPr'))
    if rp is not None:
        for child in list(rp):
            if child.tag in [qn('w:rFonts'), qn('w:sz'), qn('w:szCs'), qn('w:color')]:
                rp.remove(child)
        fonts = OxmlElement('w:rFonts')
        for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + k), 'Times New Roman')
        rp.append(fonts)
        for name, value in [('w:sz', '24'), ('w:szCs', '24'), ('w:color', '000000')]:
            e = OxmlElement(name)
            e.set(qn('w:val'), value)
            rp.append(e)

# Normalisasi final semua paragraf dan run aktif, termasuk drawing dan field.
for part in [DOC._element, DOC.sections[0].header._element, DOC.sections[0].footer._element]:
    for pe in part.xpath('.//w:p'):
        prop = pe.get_or_add_pPr()
        spacing = prop.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            prop.append(spacing)
        spacing.set(qn('w:line'), '360')
        spacing.set(qn('w:lineRule'), 'auto')
        code_style = prop.find(qn('w:pStyle'))
        family = 'Consolas' if code_style is not None and code_style.get(qn('w:val')) == cs.style_id else 'Times New Roman'
        from docx.text.run import Run
        for relem in pe.xpath('./w:r'):
            font(Run(relem, None), family=family)

memory = BytesIO()
DOC.save(memory)
memory.seek(0)
editable = {'word/document.xml', 'word/styles.xml', 'word/settings.xml', 'word/header1.xml',
            'word/footer1.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml', 'docProps/core.xml'}
with ZipFile(REFERENCE) as original, ZipFile(memory) as built, ZipFile(OUT, 'w', ZIP_DEFLATED) as final:
    for info in original.infolist():
        payload = built.read(info.filename) if info.filename in editable else original.read(info.filename)
        final.writestr(info, payload)
    for name in set(built.namelist()) - set(original.namelist()):
        final.writestr(name, built.read(name))

with ZipFile(REFERENCE) as original, ZipFile(OUT) as final:
    preserve = set(original.namelist()) - editable
    assert all(original.read(n) == final.read(n) for n in preserve)
    final_xml = etree.fromstring(final.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    texts = ''.join(final_xml.xpath('//w:t/text()', namespaces=ns))
    for name in CREATES:
        assert 'CREATE TABLE ' + name + ' (' in texts
        assert 'INSERT INTO ' + name in texts
    for part in ['word/document.xml', 'word/header1.xml', 'word/footer1.xml']:
        x = etree.fromstring(final.read(part))
        for prop in x.xpath('//w:r[w:t]/w:rPr', namespaces=ns):
            assert prop.find(qn('w:sz')).get(qn('w:val')) == '24'
            assert prop.find(qn('w:color')).get(qn('w:val')) == '000000'
        for spacing in x.xpath('//w:p/w:pPr/w:spacing', namespaces=ns):
            assert spacing.get(qn('w:line')) == '360'
            assert spacing.get(qn('w:lineRule')) == 'auto'

(TMP / 'bookmarks.json').write_text(json.dumps(BOOKMARKS, indent=2))
print(json.dumps({'docx': str(OUT), 'tables': len(DOC.tables),
                  'preserve_parts_unchanged': len(preserve), 'bookmarks': BOOKMARKS}, indent=2))
